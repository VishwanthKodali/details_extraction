from fastapi import APIRouter, UploadFile, File, HTTPException
from fastapi.responses import StreamingResponse
import openpyxl
from io import BytesIO
from app.v1.core.pdf_extraction import extract_invoice_data
from app.v1.models.detail_extraction import InvoiceResponse

router = APIRouter(prefix="/detail", tags=["extraction"])

@router.post("/extract", response_model=InvoiceResponse)
async def extract_details(file: UploadFile = File(...)):
    if not file.filename.lower().endswith('.pdf'):
        raise HTTPException(400, "PDF only")
    content = await file.read()
    data = await extract_invoice_data(content)
    return InvoiceResponse(**data)  # Auto-validates

@router.post("/to-excel")
async def export_excel(file: UploadFile = File(...)):
    content = await file.read()
    data = await extract_invoice_data(content)
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Invoice Data"
    ws['A1'], ws['B1'] = "Field", "Value"
    ws['A2'], ws['B2'] = "Invoice Number", data.get('invoice_number', '')
    ws['A3'], ws['B3'] = "Total", data.get('total', '')
    
    # Line items
    for i, item in enumerate(data.get('line_items', []), 5):
        ws[f'A{i}'] = item['desc']
        ws[f'B{i}'], ws[f'C{i}'], ws[f'D{i}'] = item['qty'], item['price'], item['total']
    
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename={file.filename[:-4]}.xlsx"})
